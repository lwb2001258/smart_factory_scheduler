from pathlib import Path
import json

import matplotlib.pyplot as plt
import numpy as np
from pptx import Presentation


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "Wenbin_Lin_MSc_Dissertation_Presentation.pptx"
OUTPUT = ROOT / "Wenbin_Lin_MSc_Dissertation_Presentation_2026_Final.pptx"
ASSETS = ROOT / "docs" / "presentation_assets_2026"
ASSETS.mkdir(parents=True, exist_ok=True)

ALGORITHMS = ["NearestNeighbour", "Hungarian", "GA", "SA", "SARSA", "DQN", "PPO_RL"]
DISPLAY = {"NearestNeighbour": "NN", "PPO_RL": "PPO"}
COLORS = {
    "NearestNeighbour": "#6B7280", "Hungarian": "#2563EB",
    "GA": "#8B5CF6", "SA": "#F59E0B", "SARSA": "#10B981",
    "DQN": "#06B6D4", "PPO_RL": "#EF4444",
}


def measured_results():
    rows = {}
    for path in (ROOT / "results").glob("experiment_?_*.json"):
        try:
            data = json.loads(path.read_text(encoding="utf-8-sig"))
        except (OSError, json.JSONDecodeError):
            continue
        synthetic = data.get("synthetic_adjustment", {})
        if synthetic.get("classification", "measured") != "measured":
            continue
        info = data.get("experiment_info", {})
        metrics = data.get("summary_metrics", {})
        scene, algorithm = info.get("scenario"), info.get("scheduler")
        if scene not in "ABC" or algorithm not in ALGORITHMS:
            continue
        completed = metrics.get("total_tasks_completed", 0)
        previous = rows.get((scene, algorithm))
        if previous is None or completed > previous["completed"]:
            rows[(scene, algorithm)] = {
                "throughput": metrics.get("throughput_per_minute", 0.0),
                "completed": completed,
                "generated": metrics.get("total_tasks_generated", 0),
                "completion_time": metrics.get("avg_task_completion_time", 0.0),
                "latency": metrics.get("scheduling_latency_p95_ms", 0.0),
                "violations": metrics.get("pair_distance_violations", 0),
            }
    return rows


def chart_style():
    plt.rcParams.update({
        "font.family": "DejaVu Sans", "font.size": 10,
        "axes.titlesize": 14, "axes.labelsize": 10,
        "axes.edgecolor": "#9CA3AF", "axes.linewidth": 0.8,
        "figure.facecolor": "white", "axes.facecolor": "white",
    })


def save_throughput(rows):
    scenes = list("ABC")
    x = np.arange(3)
    width = 0.115
    fig, ax = plt.subplots(figsize=(12.1, 3.85), dpi=170)
    for idx, alg in enumerate(ALGORITHMS):
        vals = [rows[(s, alg)]["throughput"] for s in scenes]
        ax.bar(x + (idx - 3) * width, vals, width, label=DISPLAY.get(alg, alg),
               color=COLORS[alg])
    ax.set_xticks(x, ["A · 3 robots", "B · 5 robots", "C · 8 robots"])
    ax.set_ylabel("Completed tasks per minute")
    ax.set_ylim(0, 5.8)
    ax.grid(axis="y", alpha=.18)
    ax.spines[["top", "right"]].set_visible(False)
    ax.legend(ncol=7, loc="upper left", frameon=False, fontsize=8.5)
    fig.tight_layout()
    path = ASSETS / "selected_algorithms_throughput.png"
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)
    return path


def save_scatter(rows):
    fig, ax = plt.subplots(figsize=(12.1, 3.85), dpi=170)
    for alg in ALGORITHMS:
        r = rows[("C", alg)]
        ax.scatter(r["completion_time"], r["throughput"], s=150,
                   color=COLORS[alg], edgecolor="white", linewidth=1.5, zorder=3)
        ax.annotate(DISPLAY.get(alg, alg), (r["completion_time"], r["throughput"]),
                    xytext=(5, 6), textcoords="offset points", fontsize=9)
    ax.axhspan(4.8, 5.3, color="#DCFCE7", alpha=.7)
    ax.set_xlabel("Mean task completion time (s)  ← better")
    ax.set_ylabel("Throughput (tasks/min)  → better")
    ax.set_title("Scene C: operational efficiency under high demand", loc="left")
    ax.grid(alpha=.18)
    ax.spines[["top", "right"]].set_visible(False)
    fig.tight_layout()
    path = ASSETS / "scene_c_efficiency_scatter.png"
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)
    return path


def save_rl_scaling(rows):
    scenes = list("ABC")
    best = {s: max(rows[(s, a)]["throughput"] for a in ALGORITHMS) for s in scenes}
    fig, ax = plt.subplots(figsize=(7.6, 4.55), dpi=170)
    for alg in ["SARSA", "DQN", "PPO_RL"]:
        ratios = [100 * rows[(s, alg)]["throughput"] / best[s] for s in scenes]
        ax.plot([3, 5, 8], ratios, marker="o", markersize=8, linewidth=2.7,
                color=COLORS[alg], label=DISPLAY.get(alg, alg))
        for x, y in zip([3, 5, 8], ratios):
            ax.text(x, y + 1.6, f"{y:.1f}%", ha="center", fontsize=8)
    ax.set_xticks([3, 5, 8])
    ax.set_ylim(45, 104)
    ax.set_xlabel("Fleet size (robots)")
    ax.set_ylabel("Throughput relative to best measured method")
    ax.grid(alpha=.18)
    ax.spines[["top", "right"]].set_visible(False)
    ax.legend(frameon=False, loc="center left")
    fig.tight_layout()
    path = ASSETS / "rl_relative_scaling.png"
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)
    return path


def save_heatmap(rows):
    algs = ["NearestNeighbour", "Hungarian", "GA", "SA", "SARSA", "DQN"]
    raw = []
    for a in algs:
        r = rows[("C", a)]
        raw.append([r["throughput"], r["completed"] / r["generated"],
                    r["completion_time"], r["latency"], r["violations"]])
    raw = np.asarray(raw, dtype=float)
    score = np.zeros_like(raw)
    for j in range(raw.shape[1]):
        lo, hi = raw[:, j].min(), raw[:, j].max()
        norm = (raw[:, j] - lo) / max(hi - lo, 1e-9)
        score[:, j] = norm if j < 2 else 1 - norm
    fig, ax = plt.subplots(figsize=(7.75, 4.85), dpi=170)
    im = ax.imshow(score, cmap="YlGnBu", vmin=0, vmax=1, aspect="auto")
    ax.set_xticks(range(5), ["Throughput", "Completion\nrate", "Completion\ntime", "P95 decision\nlatency", "Separation\nevents"])
    ax.set_yticks(range(len(algs)), [DISPLAY.get(a, a) for a in algs])
    for i in range(score.shape[0]):
        for j in range(score.shape[1]):
            ax.text(j, i, f"{score[i,j]:.2f}", ha="center", va="center",
                    color="white" if score[i,j] > .58 else "#111827", fontsize=8.5)
    ax.set_title("Scene C normalised score · higher is better", loc="left")
    fig.colorbar(im, ax=ax, fraction=.025, pad=.025)
    fig.tight_layout()
    path = ASSETS / "scene_c_selected_heatmap.png"
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)
    return path


def replace_picture(slide, index, image_path):
    old = slide.shapes[index]
    left, top, width, height = old.left, old.top, old.width, old.height
    old._element.getparent().remove(old._element)
    slide.shapes.add_picture(str(image_path), left, top, width, height)


def set_shape_text(shape, value):
    tf = shape.text_frame
    if tf.paragraphs and tf.paragraphs[0].runs:
        tf.paragraphs[0].runs[0].text = value
        for run in tf.paragraphs[0].runs[1:]:
            run.text = ""
        for paragraph in tf.paragraphs[1:]:
            for run in paragraph.runs:
                run.text = ""
    else:
        shape.text = value


def build():
    chart_style()
    rows = measured_results()
    required = {(s, a) for s in "ABC" for a in ALGORITHMS}
    missing = required - set(rows)
    if missing:
        raise RuntimeError(f"Missing measured results: {sorted(missing)}")
    charts = [save_throughput(rows), save_scatter(rows),
              save_rl_scaling(rows), save_heatmap(rows)]

    prs = Presentation(SOURCE)
    # Slide 12: selected-algorithm throughput comparison.
    s = prs.slides[11]
    set_shape_text(s.shapes[0], "Measured throughput across factory loads")
    set_shape_text(s.shapes[1], "SELECTED STRONG BASELINES AND LEARNING METHODS · 30-MINUTE WEBOTS RUNS")
    set_shape_text(s.shapes[6], "Scene A")
    set_shape_text(s.shapes[7], "DQN: 1.53 vs SA: 1.57")
    set_shape_text(s.shapes[8], "97.9% of the measured best")
    set_shape_text(s.shapes[10], "Scene B")
    set_shape_text(s.shapes[11], "SARSA: 2.70 vs SA: 2.80")
    set_shape_text(s.shapes[12], "96.4% of the measured best")
    set_shape_text(s.shapes[14], "Scene C")
    set_shape_text(s.shapes[15], "SARSA: 5.07 vs Hungarian: 5.20")
    set_shape_text(s.shapes[16], "97.4% of the measured best")
    set_shape_text(s.shapes[17], "Point estimates only: one measured run per algorithm–scene pair; synthetic PPO records are excluded.")
    replace_picture(s, 4, charts[0])

    # Slide 13: no minimum-wait comparison; use throughput/completion-time trade-off.
    s = prs.slides[12]
    set_shape_text(s.shapes[0], "High-density Scene C: completion efficiency")
    set_shape_text(s.shapes[1], "THROUGHPUT AND END-TO-END COMPLETION TIME · UPPER-LEFT IS THE TARGET REGION")
    set_shape_text(s.shapes[7], "SARSA")
    set_shape_text(s.shapes[8], "5.07 tasks/min: within 2.6% of the best measured throughput.")
    set_shape_text(s.shapes[11], "DQN")
    set_shape_text(s.shapes[12], "4.93 tasks/min: within 5.1%, while learning a reusable policy offline.")
    set_shape_text(s.shapes[15], "Interpretation")
    set_shape_text(s.shapes[16], "RL is competitive, but current evidence does not establish superiority over search-based methods.")
    replace_picture(s, 4, charts[1])

    # Slide 14: RL scalability relative to the best observed method.
    s = prs.slides[13]
    set_shape_text(s.shapes[8], "Competitive performance")
    set_shape_text(s.shapes[9], "SARSA remains close to the best measured method as the fleet grows: 91.5%, 96.4%, then 97.4%.")
    set_shape_text(s.shapes[12], "Learning advantage")
    set_shape_text(s.shapes[13], "Training moves search effort offline; masked inference can reuse congestion, battery and task features online.")
    set_shape_text(s.shapes[16], "Coordination boundary")
    set_shape_text(s.shapes[17], "All schedulers share the same A*, reservations, DWA and recovery stack; gains therefore arise mainly from allocation.")
    set_shape_text(s.shapes[18], "RL closes the throughput gap, but multi-seed evidence is still required.")
    replace_picture(s, 5, charts[2])

    # Slide 15: selected multi-metric heatmap without waiting-time metric.
    s = prs.slides[14]
    set_shape_text(s.shapes[0], "Selected-method comparison in Scene C")
    set_shape_text(s.shapes[1], "NORMALISED MEASURED METRICS · WAITING TIME EXCLUDED · 1 MEANS BETTER WITHIN THIS DATASET")
    set_shape_text(s.shapes[6], "Throughput")
    set_shape_text(s.shapes[7], "RL is competitive")
    set_shape_text(s.shapes[8], "SARSA reaches 97.4% and DQN 94.9% of the best.")
    set_shape_text(s.shapes[10], "Adaptation")
    set_shape_text(s.shapes[11], "Potential advantage")
    set_shape_text(s.shapes[12], "State features support battery-, queue- and congestion-aware decisions.")
    set_shape_text(s.shapes[14], "Deployment")
    set_shape_text(s.shapes[15], "Safe by design")
    set_shape_text(s.shapes[16], "Action masking, validation and deterministic fallback protect execution.")
    set_shape_text(s.shapes[18], "Evidence")
    set_shape_text(s.shapes[19], "Not yet conclusive")
    set_shape_text(s.shapes[20], "A single seed cannot prove a general learning advantage.")
    set_shape_text(s.shapes[22], "PPO")
    set_shape_text(s.shapes[23], "Redesigned")
    set_shape_text(s.shapes[24], "Pairwise PPO now shares the 278-state/161-action environment; full re-evaluation remains.")
    set_shape_text(s.shapes[25], "Completion time, decision latency and separation events: lower is better.")
    replace_picture(s, 4, charts[3])

    # Slide 16: technically grounded explanation and improvement route.
    s = prs.slides[15]
    set_shape_text(s.shapes[4],
        "●  The project implements a complete allocation-to-motion Webots loop.\n"
        "●  SA leads Scenes A/B and Hungarian leads C in the current point estimates.\n"
        "●  SARSA and DQN approach the best throughput, particularly under high load.\n"
        "●  RL has not yet surpassed online search because training abstracts physical delay and the reward only approximates final throughput.")
    set_shape_text(s.shapes[8], "One measured run per cell; seed metadata and confidence intervals are unavailable.")
    set_shape_text(s.shapes[12], "Run ≥5 matched seeds; align checkpoint selection with Webots throughput; add reward and coordination ablations.")
    set_shape_text(s.shapes[16], "Domain-randomised motion time, graph/GNN policies, charging-aware MARL, and ROS 2 hardware validation.")

    prs.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
